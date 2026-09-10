import csv
import html as htmllib
import io
import os
import re
import subprocess
import tempfile
import time
from pathlib import Path

import requests
from openpyxl import Workbook

ROOT = Path.cwd()
EXE = ROOT / "_gd_build" / "GD-Fiscal-Saude-MVP-FINAL.exe"
BASE = "http://127.0.0.1:17654"
TMP = Path(tempfile.mkdtemp(prefix="gd-fiscal-review-states-"))
DATA_DIR = TMP / "data"
ENV = os.environ.copy()
ENV["GD_FISCAL_SAUDE_TEST_MODE"] = "1"
ENV["GD_FISCAL_SAUDE_DATA_DIR"] = str(DATA_DIR)
ENV["GD_FISCAL_SAUDE_TEST_BROWSER_CAPTURE"] = str(TMP / "browser.txt")
PROCS = []


def valid_cpf(seed: int) -> str:
    base = [int(c) for c in f"{100000000 + seed:09d}"]
    s = sum(d * w for d, w in zip(base, range(10, 1, -1)))
    d1 = 11 - (s % 11)
    d1 = 0 if d1 >= 10 else d1
    a = base + [d1]
    s = sum(d * w for d, w in zip(a, range(11, 1, -1)))
    d2 = 11 - (s % 11)
    d2 = 0 if d2 >= 10 else d2
    return "".join(map(str, a + [d2]))


def workbook(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan1"
    ws["A1"] = "Carnê Leão - Livro Caixa - 2026"
    header = 84
    ws.cell(header - 1, 1, "JULHO")
    ws.cell(header, 1, "Nome")
    ws.cell(header, 7, "CPF - Responsável")
    ws.cell(header, 10, "CPF e Nome - Paciente - Se necessário")
    ws.cell(header, 16, "Data")
    ws.cell(header, 18, "Valor")
    for i, amount in enumerate([100.00, 200.00, 300.00, 400.00], start=1):
        r = header + i
        ws.cell(r, 1, f"Pagador Revisão {i}")
        ws.cell(r, 7, valid_cpf(900 + i))
        ws.cell(r, 16, i)
        ws.cell(r, 18, amount)
    wb.save(path)


def start():
    p = subprocess.Popen([str(EXE)], env=ENV)
    PROCS.append(p)
    deadline = time.time() + 30
    while time.time() < deadline:
        if p.poll() is not None:
            raise RuntimeError(f"executable exited: {p.returncode}")
        try:
            if requests.get(BASE + "/health", timeout=0.5).status_code == 200:
                return p
        except Exception:
            pass
        time.sleep(0.25)
    raise RuntimeError("health timeout")


def stop(p):
    if p.poll() is None:
        p.terminate()
        try:
            p.wait(5)
        except Exception:
            p.kill()
            p.wait(5)


def csrf(session, path="/payments"):
    body = session.get(BASE + path, timeout=5).text
    m = re.search(r'name="csrf" value="([^"]+)"', body)
    if not m:
        raise AssertionError(f"csrf not found at {path}")
    return m.group(1), body


def add_professional(session):
    token, _ = csrf(session, "/professionals")
    name = "Profissional Revisão"
    r = session.post(
        BASE + "/professionals/new",
        data={
            "csrf": token,
            "name": name,
            "cpf": "52998224725",
            "occupation_code": "255",
            "registry": "07/REV01",
        },
        allow_redirects=True,
        timeout=5,
    )
    assert r.status_code == 200
    pattern = rf"<tr><td>{re.escape(htmllib.escape(name))}</td>.*?/professionals/edit\?id=(\d+)"
    m = re.search(pattern, r.text, re.S)
    assert m, r.text[:1600]
    return int(m.group(1))


def import_book(session, xlsx: Path, pid: int):
    token, _ = csrf(session, "/import")
    with xlsx.open("rb") as f:
        preview = session.post(
            BASE + "/import/preview",
            data={"csrf": token},
            files={"spreadsheet": (xlsx.name, f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            timeout=20,
        )
    assert preview.status_code == 200
    assert "<b>4</b>importáveis" in preview.text, preview.text[:1600]
    pt = re.search(r'name="preview_token" value="([^"]+)"', preview.text)
    ct = re.search(r'name="csrf" value="([^"]+)"', preview.text)
    assert pt and ct
    commit = session.post(
        BASE + "/import/commit",
        data={"csrf": ct.group(1), "preview_token": pt.group(1), "professional_id": str(pid)},
        allow_redirects=True,
        timeout=10,
    )
    assert commit.status_code == 200
    assert "4 lançamentos novos" in commit.text


def payment_id(body: str, payer: str) -> int:
    pattern = rf'<tr data-payment-id="(\d+)">(?:(?!</tr>).)*{re.escape(payer)}(?:(?!</tr>).)*</tr>'
    m = re.search(pattern, body, re.S)
    if not m:
        raise AssertionError(f"payment id not found for {payer}; page prefix={body[:1800]!r}")
    return int(m.group(1))


def review(session, pid: int, status: str, reason: str = ""):
    token, _ = csrf(session, "/payments")
    return session.post(
        BASE + "/payments/review",
        data={
            "csrf": token,
            "payment_id": str(pid),
            "status": status,
            "non_launch_reason": reason,
        },
        allow_redirects=True,
        timeout=5,
    )


def filter_status(session, professional_id: int, status: str):
    return session.get(
        BASE + f"/payments?professional_id={professional_id}&year=2026&month=7&status={status}",
        timeout=5,
    ).text


def export_rows(session, professional_id: int):
    token, _ = csrf(session, "/export")
    r = session.post(
        BASE + "/export",
        data={"csrf": token, "professional_id": str(professional_id), "year": "2026", "month": "7"},
        allow_redirects=False,
        timeout=10,
    )
    assert r.status_code == 200, (r.status_code, r.text[:1200])
    return list(csv.reader(io.StringIO(r.content.decode("utf-8")), delimiter=";"))


try:
    p = start()
    s = requests.Session()
    professional_id = add_professional(s)
    xlsx = TMP / "review.xlsx"
    workbook(xlsx)
    import_book(s, xlsx, professional_id)

    page = filter_status(s, professional_id, "")
    assert "Validado para Receita Saúde" in page
    assert "Validado para Nota Fiscal" in page
    assert "Não lançado" in page
    assert "Confirmar lançados no Receita Saúde" in s.get(BASE + "/export", timeout=5).text

    ids = {i: payment_id(page, f"Pagador Revisão {i}") for i in range(1, 5)}

    # Não lançado exige justificativa e não altera o item se o motivo faltar.
    bad = review(s, ids[3], "NAO_LANCADO", "")
    assert "motivo" in bad.text.lower()
    assert "Pagador Revisão 3" in filter_status(s, professional_id, "PENDENTE_REVISAO")

    assert review(s, ids[1], "VALIDADO_RECEITA_SAUDE").status_code == 200
    assert review(s, ids[2], "VALIDADO_NOTA_FISCAL").status_code == 200
    assert review(s, ids[3], "NAO_LANCADO", "Pagamento não sujeito a lançamento neste período").status_code == 200

    # O comando em lote só alcança os ainda pendentes.
    token, _ = csrf(s, "/payments")
    bulk = s.post(
        BASE + "/payments/validate",
        data={"csrf": token, "professional_id": str(professional_id), "year": "2026", "month": "7"},
        allow_redirects=True,
        timeout=5,
    )
    assert "1 lançamentos validados" in bulk.text, bulk.text[:1200]

    nota = filter_status(s, professional_id, "VALIDADO_NOTA_FISCAL")
    assert "Pagador Revisão 2" in nota and "Pagador Revisão 1" not in nota and "Pagador Revisão 4" not in nota
    nao = filter_status(s, professional_id, "NAO_LANCADO")
    assert "Pagador Revisão 3" in nao
    assert "Pagamento não sujeito a lançamento neste período" in nao
    receita = filter_status(s, professional_id, "VALIDADO_RECEITA_SAUDE")
    assert "Pagador Revisão 1" in receita and "Pagador Revisão 4" in receita

    # Exportar não sela nem muda o estado; somente os validados para Receita Saúde entram no CSV.
    rows = export_rows(s, professional_id)
    assert len(rows) == 2 and all(len(row) == 16 for row in rows)
    receita_after_export = filter_status(s, professional_id, "VALIDADO_RECEITA_SAUDE")
    assert "Pagador Revisão 1" in receita_after_export and "Pagador Revisão 4" in receita_after_export

    # A decisão pode ser corrigida livremente antes da confirmação final.
    assert review(s, ids[1], "VALIDADO_NOTA_FISCAL").status_code == 200
    assert "Pagador Revisão 1" in filter_status(s, professional_id, "VALIDADO_NOTA_FISCAL")
    assert "Pagador Revisão 1" not in filter_status(s, professional_id, "VALIDADO_RECEITA_SAUDE")
    assert review(s, ids[1], "VALIDADO_RECEITA_SAUDE").status_code == 200

    # Confirmação explícita sela apenas os atuais VALIDADO_RECEITA_SAUDE.
    token, _ = csrf(s, "/export")
    confirmed = s.post(
        BASE + "/payments/confirm-receita",
        data={"csrf": token, "professional_id": str(professional_id), "year": "2026", "month": "7"},
        allow_redirects=True,
        timeout=5,
    )
    assert "2 lançamentos confirmados" in confirmed.text, confirmed.text[:1200]
    final = filter_status(s, professional_id, "LANCADO_RECEITA_SAUDE")
    assert "Pagador Revisão 1" in final and "Pagador Revisão 4" in final
    assert "Lançado no Receita Saúde" in final

    # Depois de selado, o item não pode voltar para outro estado por edição individual.
    locked = review(s, ids[1], "PENDENTE_REVISAO")
    assert "já foi confirmado" in locked.text.lower()
    final_again = filter_status(s, professional_id, "LANCADO_RECEITA_SAUDE")
    assert "Pagador Revisão 1" in final_again

    # Estados que não eram Receita permanecem intactos.
    assert "Pagador Revisão 2" in filter_status(s, professional_id, "VALIDADO_NOTA_FISCAL")
    assert "Pagador Revisão 3" in filter_status(s, professional_id, "NAO_LANCADO")

    stop(p)
    p = start()
    s = requests.Session()
    assert "Pagador Revisão 1" in filter_status(s, professional_id, "LANCADO_RECEITA_SAUDE")
    assert "Pagamento não sujeito a lançamento neste período" in filter_status(s, professional_id, "NAO_LANCADO")

    print("REVIEW_STATES_GATE_GREEN")
finally:
    for proc in PROCS:
        stop(proc)
