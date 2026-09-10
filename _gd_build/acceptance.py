import csv, io, os, re, shutil, subprocess, tempfile, time
from pathlib import Path
import requests
from openpyxl import Workbook

ROOT = Path.cwd()
EXE = ROOT / '_gd_build' / 'GD-Fiscal-Saude-MVP-FINAL.exe'
BASE = 'http://127.0.0.1:17654'
td = Path(tempfile.mkdtemp(prefix='gd-fiscal-mvp-'))
data_dir = td / 'data'
capture = td / 'browser.txt'
env = os.environ.copy()
env['GD_FISCAL_SAUDE_TEST_MODE'] = '1'
env['GD_FISCAL_SAUDE_DATA_DIR'] = str(data_dir)
env['GD_FISCAL_SAUDE_TEST_BROWSER_CAPTURE'] = str(capture)
procs = []

def valid_cpf(seed):
    base = [int(c) for c in f'{100000000 + seed:09d}']
    s = sum(d*w for d,w in zip(base, range(10,1,-1)))
    d1 = 11-(s%11); d1 = 0 if d1 >= 10 else d1
    a = base + [d1]
    s = sum(d*w for d,w in zip(a, range(11,1,-1)))
    d2 = 11-(s%11); d2 = 0 if d2 >= 10 else d2
    return ''.join(map(str, a+[d2]))

def workbook(path, month, rows, blank=0, invalid=False):
    wb = Workbook(); ws = wb.active; ws.title = 'Plan1'
    ws['A1'] = 'Carnê Leão - Livro Caixa - 2026'
    month_names = {7:'JULHO',8:'AGOSTO',9:'SETEMBRO'}; header = 84
    ws.cell(header-1,1,month_names[month])
    ws.cell(header,1,'Nome'); ws.cell(header,7,'CPF - Responsável')
    ws.cell(header,10,'CPF e Nome - Paciente - Se necessário')
    ws.cell(header,16,'Data'); ws.cell(header,18,'Valor')
    amounts = [400.0]*25 + [870.0] if month == 7 and rows == 26 else [123.45+i for i in range(rows)]
    for i, amount in enumerate(amounts):
        r = header+1+i
        ws.cell(r,1,f'Pagador M{month} {i+1}')
        ws.cell(r,7,valid_cpf(month*100+i+1))
        ws.cell(r,16,i+1); ws.cell(r,18,amount)
    for i in range(blank):
        r = header+1+rows+i
        ws.cell(r,1,f'Sem atendimento {i+1}')
        ws.cell(r,7,valid_cpf(month*100+rows+i+1))
    if invalid:
        r = header+1+rows+blank
        ws.cell(r,1,'CPF inválido'); ws.cell(r,7,'12345678901'); ws.cell(r,16,20); ws.cell(r,18,250)
    tr = header+1+rows+blank+(1 if invalid else 0)
    ws.cell(tr,1,'Valor Total Mês'); ws.cell(tr,18,sum(amounts))
    wb.save(path)

def start(exe=EXE):
    p = subprocess.Popen([str(exe)], env=env); procs.append(p)
    deadline = time.time()+30
    while time.time() < deadline:
        if p.poll() is not None: raise RuntimeError(f'executable exited: {p.returncode}')
        try:
            if requests.get(BASE+'/health', timeout=.5).status_code == 200: return p
        except Exception: pass
        time.sleep(.25)
    raise RuntimeError('health timeout')

def stop(p):
    if p.poll() is None:
        p.terminate()
        try: p.wait(5)
        except Exception: p.kill(); p.wait(5)

def csrf(session, path='/professionals'):
    h = session.get(BASE+path, timeout=5).text
    m = re.search(r'name="csrf" value="([^"]+)"', h)
    if not m: raise AssertionError('csrf not found')
    return m.group(1), h

def professional_id(html, name):
    m = re.search(r'<option value="(\d+)"[^>]*>'+re.escape(name)+r'\s+—', html)
    if not m: raise AssertionError(f'professional id not found for {name}')
    return int(m.group(1))

def import_book(session, xlsx, pid, expect_valid, expect_skipped, expect_restrict=0):
    token, _ = csrf(session, '/import')
    with open(xlsx, 'rb') as f:
        resp = session.post(BASE+'/import/preview', data={'csrf':token}, files={'spreadsheet':(xlsx.name,f,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}, timeout=20)
    assert resp.status_code == 200, resp.status_code
    h = resp.text
    assert f'<b>{expect_valid}</b>importáveis' in h, h[:1200]
    assert f'<b>{expect_skipped}</b>sem atendimento' in h
    assert f'<b>{expect_restrict}</b>pendências reais' in h
    m = re.search(r'name="preview_token" value="([^"]+)"', h); assert m
    token2 = re.search(r'name="csrf" value="([^"]+)"', h).group(1)
    rr = session.post(BASE+'/import/commit', data={'csrf':token2,'preview_token':m.group(1),'professional_id':str(pid)}, allow_redirects=True, timeout=10)
    assert rr.status_code == 200
    return rr.text

try:
    p = start(); s = requests.Session()
    token, _ = csrf(s)
    A = 'Profissional A'; ACPF = '52998224725'
    r = s.post(BASE+'/professionals/new', data={'csrf':token,'name':A,'cpf':ACPF,'occupation_code':'255','registry':'07/54321'}, allow_redirects=True, timeout=5)
    assert r.status_code == 200 and A in r.text
    _, imp = csrf(s, '/import'); aid = professional_id(imp, A)

    a = td/'julho.xlsx'; workbook(a, 7, 26, blank=4)
    h = import_book(s, a, aid, 26, 4, 0)
    assert '26 lançamentos novos' in h and '4 sem atendimento' in h
    pay = s.get(BASE+f'/payments?professional_id={aid}&year=2026&month=7', timeout=5).text
    assert 'Pagador M7 1' in pay and 'Pagador M7 26' in pay and 'R$ 10.870,00' in pay

    h = import_book(s, a, aid, 26, 4, 0)
    assert '0 lançamentos novos' in h and '26 duplicados ignorados' in h

    bad = td/'pendencia.xlsx'; workbook(bad, 9, 0, invalid=True)
    h = import_book(s, bad, aid, 0, 0, 1)
    assert '1 pendências' in h
    restr = s.get(BASE+'/restrictions', timeout=5).text
    assert 'CPF_INVALIDO' in restr and A in restr

    stop(p); p = start(); s = requests.Session()
    assert A in s.get(BASE+'/professionals', timeout=5).text
    assert 'Pagador M7 26' in s.get(BASE+f'/payments?professional_id={aid}&year=2026&month=7', timeout=5).text

    token, _ = csrf(s)
    B = 'Profissional B'; BCPF = '11144477735'
    r = s.post(BASE+'/professionals/new', data={'csrf':token,'name':B,'cpf':BCPF,'occupation_code':'255','registry':'07/99999'}, allow_redirects=True, timeout=5)
    assert r.status_code == 200 and B in r.text
    _, imp = csrf(s, '/import'); bid = professional_id(imp, B)
    b = td/'agosto.xlsx'; workbook(b, 8, 2)
    h = import_book(s, b, bid, 2, 0, 0); assert '2 lançamentos novos' in h
    aview = s.get(BASE+f'/payments?professional_id={aid}&year=2026&month=7', timeout=5).text
    bview = s.get(BASE+f'/payments?professional_id={bid}&year=2026&month=8', timeout=5).text
    assert 'Pagador M8' not in aview and 'Pagador M7' not in bview

    token, _ = csrf(s, '/payments')
    rr = s.post(BASE+'/payments/validate', data={'csrf':token,'professional_id':str(aid),'year':'2026','month':'7'}, allow_redirects=True, timeout=5)
    assert '26 lançamentos validados' in rr.text
    token, _ = csrf(s, '/export')
    out = s.post(BASE+'/export', data={'csrf':token,'professional_id':str(aid),'year':'2026','month':'7'}, allow_redirects=False, timeout=10)
    assert out.status_code == 200
    text = out.content.decode('utf-8'); lines = [x for x in text.splitlines() if x.strip()]
    parsed = list(csv.reader(io.StringIO(text), delimiter=';'))
    assert len(lines) == 26 and all(len(x)==16 for x in parsed) and all(x[14]==ACPF for x in parsed)
    total = sum(float(x[3].replace('.','').replace(',','.')) for x in parsed)
    assert abs(total-10870) < 0.001

    token, _ = csrf(s, '/payments')
    rr = s.post(BASE+'/payments/validate', data={'csrf':token,'professional_id':str(bid),'year':'2026','month':'8'}, allow_redirects=True, timeout=5)
    assert '2 lançamentos validados' in rr.text
    token, _ = csrf(s, '/export')
    outb = s.post(BASE+'/export', data={'csrf':token,'professional_id':str(bid),'year':'2026','month':'8'}, allow_redirects=False, timeout=10)
    rowsb = list(csv.reader(io.StringIO(outb.content.decode()), delimiter=';'))
    assert len(rowsb)==2 and all(len(x)==16 for x in rowsb) and all(x[14]==BCPF for x in rowsb) and all(x[14]!=ACPF for x in rowsb)

    stop(p)
    replacement = td/'GD-Fiscal-Saude-MVP-FINAL-reinstalled.exe'; shutil.copy2(EXE, replacement)
    p = start(replacement); s = requests.Session()
    prof = s.get(BASE+'/professionals', timeout=5).text
    assert A in prof and B in prof
    dash = s.get(BASE+'/', timeout=5).text
    assert '<b>2</b>profissionais' in dash and '<b>28</b>lançamentos salvos' in dash
    print('MVP_GATE_GREEN')
finally:
    for p in procs: stop(p)
