import html as htmllib
import os
import re
import subprocess
import tempfile
import time
from pathlib import Path

import requests
from openpyxl import Workbook

ROOT = Path.cwd()
LEGACY = ROOT / "_gd_build" / "GD-Fiscal-Saude-MVP-LEGACY.exe"
NEW = ROOT / "_gd_build" / "GD-Fiscal-Saude-MVP-FINAL.exe"
BASE = "http://127.0.0.1:17654"
TMP = Path(tempfile.mkdtemp(prefix="gd-fiscal-review-upgrade-"))
DATA_DIR = TMP / "data"
ENV = os.environ.copy()
ENV["GD_FISCAL_SAUDE_TEST_MODE"] = "1"
ENV["GD_FISCAL_SAUDE_DATA_DIR"] = str(DATA_DIR)
ENV["GD_FISCAL_SAUDE_TEST_BROWSER_CAPTURE"] = str(TMP / "browser.txt")
PROCS = []


def valid_cpf(seed: int) -> str:
    base = [int(c) for c in f"{100000000 + seed:09d}"]
    s = sum(d * w for d, w in zip(base, range(10, 1, -1)))
    d1 = 11 - (s % 11)
    d1 = 0 if d1 >= 10 else d1
    a = base + [d1]
    s = sum(d * w for d, w in zip(a, range(11, 1, -1)))
    d2 = 11 - (s % 11)
    d2 = 0 if d2 >= 10 else d2
    return "".join(map(str, a + [d2]))


def make_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan1"
    ws["A1"] = "Carnê Leão - Livro Caixa - 2026"
    header = 84
    ws.cell(header - 1, 1, "JULHO")
    ws.cell(header, 1, "Nome")
    ws.cell(header, 7, "CPF - Responsável")
    ws.cell(header, 10, "CPF e Nome - Paciente - Se necessário")
    ws.cell(header, 16, "Data")
    ws.cell(header, 18, "Valor")
    for i, amount in enumerate([150.00, 250.00], start=1):
        r = header + i
        ws.cell(r, 1, f"Pagador Upgrade {i}")
        ws.cell(r, 7, valid_cpf(980 + i))
        ws.cell(r, 16, i)
        ws.cell(r, 18, amount)
    wb.save(path)


def start(exe: Path):
    p = subprocess.Popen([str(exe)], env=ENV)
    PROCS.append(p)
    deadline = time.time() + 30
    while time.time() < deadline:
        if p.poll() is not None:
            raise RuntimeError(f"{exe.name} exited: {p.returncode}")
        try:
            if requests.get(BASE + "/health", timeout=0.5).status_code == 200:
                return p
        except Exception:
            pass
        time.sleep(0.25)
    raise RuntimeError("health timeout")


def stop(p):
    if p.poll() is None:
        p.terminate()
        try:
            p.wait(5)
        except Exception:
            p.kill()
            p.wait(5)


def csrf(s, path):
    body = s.get(BASE + path, timeout=5).text
    m = re.search(r'name="csrf" value="([^"]+)"', body)
    assert m, body[:1000]
    return m.group(1), body


def add_professional(s):
    token, _ = csrf(s, "/professionals")
    name = "Profissional Upgrade"
    r = s.post(BASE + "/professionals/new", data={"csrf": token, "name": name, "cpf": "52998224725", "occupation_code": "255", "registry": "07/UPG01"}, allow_redirects=True, timeout=5)
    assert r.status_code == 200
    m = re.search(rf"<tr><td>{re.escape(htmllib.escape(name))}</td>.*?/professionals/edit\?id=(\d+)", r.text, re.S)
    assert m, r.text[:1200]
    return int(m.group(1))


def import_book(s, path: Path, pid: int):
    token, _ = csrf(s, "/import")
    with path.open("rb") as f:
        preview = s.post(BASE + "/import/preview", data={"csrf": token}, files={"spreadsheet": (path.name, f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, timeout=20)
    assert "<b>2</b>importáveis" in preview.text
    pt = re.search(r'name="preview_token" value="([^"]+)"', preview.text)
    ct = re.search(r'name="csrf" value="([^"]+)"', preview.text)
    assert pt and ct
    r = s.post(BASE + "/import/commit", data={"csrf": ct.group(1), "preview_token": pt.group(1), "professional_id": str(pid)}, allow_redirects=True, timeout=10)
    assert "2 lançamentos novos" in r.text


try:
    legacy = start(LEGACY)
    s = requests.Session()
    pid = add_professional(s)
    book = TMP / "upgrade.xlsx"
    make_book(book)
    import_book(s, book, pid)

    token, _ = csrf(s, "/payments")
    r = s.post(BASE + "/payments/validate", data={"csrf": token, "professional_id": str(pid), "year": "2026", "month": "7"}, allow_redirects=True, timeout=5)
    assert "2 lançamentos validados" in r.text

    token, _ = csrf(s, "/export")
    exported = s.post(BASE + "/export", data={"csrf": token, "professional_id": str(pid), "year": "2026", "month": "7"}, allow_redirects=False, timeout=10)
    assert exported.status_code == 200
    legacy_exported = s.get(BASE + f"/payments?professional_id={pid}&year=2026&month=7&status=EXPORTADO_CSV", timeout=5).text
    assert "Pagador Upgrade 1" in legacy_exported and "Pagador Upgrade 2" in legacy_exported
    stop(legacy)

    upgraded = start(NEW)
    s = requests.Session()
    listing = s.get(BASE + "/professionals", timeout=5).text
    assert "Profissional Upgrade" in listing

    # Abrir a escrituração executa a migração aditiva. Dados previamente 'EXPORTADO_CSV'
    # voltam ao estado editável, pois exportar na versão antiga não comprovava lançamento efetivo.
    ready = s.get(BASE + f"/payments?professional_id={pid}&year=2026&month=7&status=VALIDADO_RECEITA_SAUDE", timeout=5).text
    assert "Pagador Upgrade 1" in ready and "Pagador Upgrade 2" in ready
    assert "Validado para Receita Saúde" in ready
    old = s.get(BASE + f"/payments?professional_id={pid}&year=2026&month=7&status=EXPORTADO_CSV", timeout=5).text
    assert "Pagador Upgrade 1" not in old and "Pagador Upgrade 2" not in old

    db = DATA_DIR / "fiscal_saude.sqlite3"
    assert db.exists() and db.stat().st_size > 0
    print("REVIEW_UPGRADE_GATE_GREEN")
finally:
    for proc in PROCS:
        stop(proc)
