import csv
import html as htmllib
import io
import os
import re
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

import requests
from openpyxl import Workbook

ROOT = Path.cwd()
EXE = ROOT / "_gd_build" / "GD-Fiscal-Saude-MVP-FINAL.exe"
BASE = "http://127.0.0.1:17654"
TMP = Path(tempfile.mkdtemp(prefix="gd-fiscal-mvp-gate-"))
DATA_DIR = TMP / "data"
CAPTURE = TMP / "browser.txt"
ENV = os.environ.copy()
ENV["GD_FISCAL_SAUDE_TEST_MODE"] = "1"
ENV["GD_FISCAL_SAUDE_DATA_DIR"] = str(DATA_DIR)
ENV["GD_FISCAL_SAUDE_TEST_BROWSER_CAPTURE"] = str(CAPTURE)
PROCS = []


def valid_cpf(seed: int) -> str:
    base = [int(c) for c in f"{100000000 + seed:09d}"]
    s = sum(d * w for d, w in zip(base, range(10, 1, -1)))
    d1 = 11 - (s % 11)
    d1 = 0 if d1 >= 10 else d1
    a = base + [d1]
    s = sum(d * w for d, w in zip(a, range(11, 1, -1)))
    d2 = 11 - (s % 11)
    d2 = 0 if d2 >= 10 else d2
    return "".join(map(str, a + [d2]))


def workbook(path: Path, month: int, rows: int, blank: int = 0, invalid: bool = False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan1"
    ws["A1"] = "Carnê Leão - Livro Caixa - 2026"
    month_names = {7: "JULHO", 8: "AGOSTO", 9: "SETEMBRO"}
    header = 84
    ws.cell(header - 1, 1, month_names[month])
    ws.cell(header, 1, "Nome")
    ws.cell(header, 7, "CPF - Responsável")
    ws.cell(header, 10, "CPF e Nome - Paciente - Se necessário")
    ws.cell(header, 16, "Data")
    ws.cell(header, 18, "Valor")
    amounts = [400.0] * 25 + [870.0] if month == 7 and rows == 26 else [123.45 + i for i in range(rows)]
    for i, amount in enumerate(amounts):
        r = header + 1 + i
        ws.cell(r, 1, f"Pagador M{month} {i + 1}")
        ws.cell(r, 7, valid_cpf(month * 100 + i + 1))
        ws.cell(r, 16, i + 1)
        ws.cell(r, 18, amount)
    for i in range(blank):
        r = header + 1 + rows + i
        ws.cell(r, 1, f"Sem atendimento {i + 1}")
        ws.cell(r, 7, valid_cpf(month * 100 + rows + i + 1))
    if invalid:
        r = header + 1 + rows + blank
        ws.cell(r, 1, "CPF inválido")
        ws.cell(r, 7, "12345678901")
        ws.cell(r, 16, 20)
        ws.cell(r, 18, 250)
    tr = header + 1 + rows + blank + (1 if invalid else 0)
    ws.cell(tr, 1, "Valor Total Mês")
    ws.cell(tr, 18, sum(amounts))
    wb.save(path)


def start(exe: Path = EXE):
    p = subprocess.Popen([str(exe)], env=ENV)
    PROCS.append(p)
    deadline = time.time() + 30
    while time.time() < deadline:
        if p.poll() is not None:
            errfile = DATA_DIR / "startup-error.txt"
            detail = errfile.read_text(errors="replace") if errfile.exists() else ""
            raise RuntimeError(f"executable exited: {p.returncode}\n{detail}")
        try:
            if requests.get(BASE + "/health", timeout=0.5).status_code == 200:
                return p
        except Exception:
            pass
        time.sleep(0.25)
    raise RuntimeError("health timeout")


def stop(p):
    if p.poll() is None:
        p.terminate()
        try:
            p.wait(5)
        except Exception:
            p.kill()
            p.wait(5)


def get_csrf(session, path="/professionals"):
    body = session.get(BASE + path, timeout=5).text
    m = re.search(r'name="csrf" value="([^"]+)"', body)
    if not m:
        raise AssertionError(f"csrf not found at {path}")
    return m.group(1), body


def professional_id_from_listing(body: str, name: str) -> int:
    name_escaped = re.escape(htmllib.escape(name))
    pattern = rf"<tr><td>{name_escaped}</td>.*?/professionals/edit\?id=(\d+)"
    m = re.search(pattern, body, re.S)
    if not m:
        raise AssertionError(f"professional id not found for {name}; listing prefix={body[:1200]!r}")
    return int(m.group(1))


def add_professional(session, name: str, cpf: str, registry: str) -> int:
    token, _ = get_csrf(session, "/professionals")
    r = session.post(
        BASE + "/professionals/new",
        data={"csrf": token, "name": name, "cpf": cpf, "occupation_code": "255", "registry": registry},
        allow_redirects=True,
        timeout=5,
    )
    assert r.status_code == 200, r.status_code
    assert name in r.text, r.text[:1200]
    return professional_id_from_listing(r.text, name)


def import_book(session, xlsx: Path, pid: int, expect_valid: int, expect_skipped: int, expect_restrict: int = 0):
    token, _ = get_csrf(session, "/import")
    with xlsx.open("rb") as f:
        resp = session.post(
            BASE + "/import/preview",
            data={"csrf": token},
            files={"spreadsheet": (xlsx.name, f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            timeout=20,
        )
    assert resp.status_code == 200, resp.status_code
    body = resp.text
    assert f"<b>{expect_valid}</b>importáveis" in body, body[:1600]
    assert f"<b>{expect_skipped}</b>sem atendimento" in body, body[:1600]
    assert f"<b>{expect_restrict}</b>pendências reais" in body, body[:1600]
    m = re.search(r'name="preview_token" value="([^"]+)"', body)
    assert m, body[:1600]
    csrf2 = re.search(r'name="csrf" value="([^"]+)"', body)
    assert csrf2
    commit = session.post(
        BASE + "/import/commit",
        data={"csrf": csrf2.group(1), "preview_token": m.group(1), "professional_id": str(pid)},
        allow_redirects=True,
        timeout=10,
    )
    assert commit.status_code == 200, commit.status_code
    return commit.text


def validate_month(session, pid: int, month: int, expected_count: int):
    token, _ = get_csrf(session, "/payments")
    r = session.post(
        BASE + "/payments/validate",
        data={"csrf": token, "professional_id": str(pid), "year": "2026", "month": str(month)},
        allow_redirects=True,
        timeout=5,
    )
    assert r.status_code == 200
    assert f"{expected_count} lançamentos validados" in r.text, r.text[:1200]


def export_month(session, pid: int, month: int):
    token, _ = get_csrf(session, "/export")
    r = session.post(
        BASE + "/export",
        data={"csrf": token, "professional_id": str(pid), "year": "2026", "month": str(month)},
        allow_redirects=False,
        timeout=10,
    )
    assert r.status_code == 200, (r.status_code, r.text[:1200])
    return list(csv.reader(io.StringIO(r.content.decode("utf-8")), delimiter=";"))


try:
    p = start()
    s = requests.Session()

    A = "Profissional A"
    ACPF = "52998224725"
    aid = add_professional(s, A, ACPF, "07/54321")

    july = TMP / "julho.xlsx"
    workbook(july, 7, 26, blank=4)
    body = import_book(s, july, aid, 26, 4, 0)
    assert "26 lançamentos novos" in body and "4 sem atendimento" in body
    pay_a = s.get(BASE + f"/payments?professional_id={aid}&year=2026&month=7", timeout=5).text
    assert "Pagador M7 1" in pay_a and "Pagador M7 26" in pay_a
    assert "R$ 10.870,00" in pay_a

    body = import_book(s, july, aid, 26, 4, 0)
    assert "0 lançamentos novos" in body and "26 duplicados ignorados" in body

    bad = TMP / "pendencia.xlsx"
    workbook(bad, 9, 0, invalid=True)
    body = import_book(s, bad, aid, 0, 0, 1)
    assert "1 pendências" in body
    restrictions = s.get(BASE + "/restrictions", timeout=5).text
    assert "CPF_INVALIDO" in restrictions and A in restrictions

    stop(p)
    p = start()
    s = requests.Session()
    listing = s.get(BASE + "/professionals", timeout=5).text
    assert A in listing
    assert "Pagador M7 26" in s.get(BASE + f"/payments?professional_id={aid}&year=2026&month=7", timeout=5).text

    B = "Profissional B"
    BCPF = "11144477735"
    bid = add_professional(s, B, BCPF, "07/99999")
    august = TMP / "agosto.xlsx"
    workbook(august, 8, 2)
    body = import_book(s, august, bid, 2, 0, 0)
    assert "2 lançamentos novos" in body

    aview = s.get(BASE + f"/payments?professional_id={aid}&year=2026&month=7", timeout=5).text
    bview = s.get(BASE + f"/payments?professional_id={bid}&year=2026&month=8", timeout=5).text
    assert "Pagador M8" not in aview
    assert "Pagador M7" not in bview

    validate_month(s, aid, 7, 26)
    rows_a = export_month(s, aid, 7)
    assert len(rows_a) == 26
    assert all(len(row) == 16 for row in rows_a)
    assert all(row[14] == ACPF for row in rows_a)
    total_a = sum(float(row[3].replace(".", "").replace(",", ".")) for row in rows_a)
    assert abs(total_a - 10870.0) < 0.001

    validate_month(s, bid, 8, 2)
    rows_b = export_month(s, bid, 8)
    assert len(rows_b) == 2
    assert all(len(row) == 16 for row in rows_b)
    assert all(row[14] == BCPF for row in rows_b)
    assert all(row[14] != ACPF for row in rows_b)

    stop(p)
    replacement = TMP / "GD-Fiscal-Saude-MVP-FINAL-reinstalled.exe"
    shutil.copy2(EXE, replacement)
    p = start(replacement)
    s = requests.Session()
    listing = s.get(BASE + "/professionals", timeout=5).text
    assert A in listing and B in listing
    dash = s.get(BASE + "/", timeout=5).text
    assert "<b>2</b>profissionais" in dash
    assert "<b>28</b>lançamentos salvos" in dash

    print("MVP_GATE_GREEN")
    print(f"data_dir={DATA_DIR}")
finally:
    for proc in PROCS:
        stop(proc)
